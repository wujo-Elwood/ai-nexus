# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


# 生成文档抽取测试文件
def generateTestDocuments():
    # 第1步：创建输出目录
    outputDir = Path("D:/workspace/wujo_rag/test-documents/extract")
    outputDir.mkdir(parents=True, exist_ok=True)

    # 第2步：准备模拟正式合同的数据
    contractList = [
        {
            "fileName": "contract_001_procurement_service",
            "title": "企业知识库平台采购服务合同",
            "contractNo": "HT-2026-ZS-0018",
            "partyA": "星河智能科技有限公司",
            "partyB": "云舟数据服务（上海）有限公司",
            "amount": "人民币 286,000.00 元（大写：贰拾捌万陆仟元整）",
            "signDate": "2026年06月15日",
            "startDate": "2026年07月01日",
            "endDate": "2027年06月30日",
            "contactPerson": "李明轩",
            "contactPhone": "13800001234",
            "scope": "乙方向甲方提供企业知识库平台部署、文档解析、向量检索、权限配置、使用培训及三个月上线支持服务。",
            "payment": "合同签署后五个工作日内支付合同总额的50%；系统验收通过后十个工作日内支付剩余50%。",
        },
        {
            "fileName": "contract_002_software_development",
            "title": "设备巡检系统软件开发委托合同",
            "contractNo": "BC-RD-2026-042",
            "partyA": "启明智能制造股份有限公司",
            "partyB": "北辰软件技术有限公司",
            "amount": "人民币 450,000.00 元（大写：肆拾伍万元整）",
            "signDate": "2026年05月28日",
            "startDate": "2026年06月10日",
            "endDate": "2026年12月31日",
            "contactPerson": "周婉晴",
            "contactPhone": "13988886666",
            "scope": "乙方负责设备巡检移动端、后台管理端、数据报表、异常预警及接口对接功能的设计、开发、测试和交付。",
            "payment": "甲方按里程碑付款：需求确认后支付30%，测试环境交付后支付40%，正式验收后支付30%。",
        },
        {
            "fileName": "contract_003_data_processing",
            "title": "数据处理与保密服务协议",
            "contractNo": "QW-DPA-2026-009",
            "partyA": "远景医疗科技集团有限公司",
            "partyB": "青梧智能信息科技有限公司",
            "amount": "人民币 128,800.00 元（大写：壹拾贰万捌仟捌佰元整）",
            "signDate": "2026年06月02日",
            "startDate": "2026年06月05日",
            "endDate": "2026年09月04日",
            "contactPerson": "陈思源",
            "contactPhone": "13755550128",
            "scope": "乙方为甲方提供非结构化文档脱敏、分类标注、结构化抽取和结果校验服务，处理数据仅限本协议约定项目使用。",
            "payment": "服务开始前支付合同总额的40%，中期交付通过后支付30%，最终交付并完成销毁确认后支付30%。",
        },
    ]

    # 第3步：准备 PDF 中文字体
    fontName = registerChineseFont()

    # 第4步：逐份生成 DOCX 和 PDF 文件
    for contractData in contractList:
        paragraphList = buildContractParagraphs(contractData)
        createDocxFile(outputDir, contractData, paragraphList)
        createPdfFile(outputDir, contractData, paragraphList, fontName)

    # 第5步：生成测试说明
    readmePath = outputDir / "README.txt"
    readmePath.write_text(
        "这些文件是文档抽取模块测试用的虚构合同样本，可直接上传到“文档抽取”页面测试 contract_basic 模板。\n"
        "建议先上传 PDF，再上传 DOCX，对比抽取字段结果和原文证据。\n",
        encoding="utf-8",
    )

    # 第6步：输出生成结果
    for filePath in sorted(outputDir.iterdir()):
        print(filePath)


# 注册 PDF 中文字体
def registerChineseFont():
    # 第1步：按 Windows 常见字体路径查找可用中文字体
    fontPathList = [
        "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simsun.ttc",
        "C:/Windows/Fonts/simhei.ttf",
    ]

    # 第2步：找到字体后注册给 ReportLab
    for fontPath in fontPathList:
        if Path(fontPath).exists():
            pdfmetrics.registerFont(TTFont("ChineseFont", fontPath))
            return "ChineseFont"

    # 第3步：没有中文字体时退回默认字体
    return "Helvetica"


# 构造合同正文段落
def buildContractParagraphs(contractData):
    # 第1步：把模板字段放入正式合同风格的条款中
    return [
        f"合同编号：{contractData['contractNo']}",
        f"甲方：{contractData['partyA']}",
        f"乙方：{contractData['partyB']}",
        "根据《中华人民共和国民法典》及相关法律法规，甲乙双方本着平等、自愿、公平和诚实信用原则，就本合同项下服务事项达成如下约定。",
        f"一、合同名称：{contractData['title']}。",
        f"二、服务范围：{contractData['scope']}",
        f"三、合同金额：{contractData['amount']}。该金额已包含乙方完成本合同约定服务所需的人工、管理、税费及其他合理成本。",
        f"四、服务期限：本合同服务开始日期为{contractData['startDate']}，结束日期为{contractData['endDate']}。如双方需延长服务期限，应另行签署书面补充协议。",
        f"五、付款方式：{contractData['payment']}",
        "六、验收标准：乙方提交的交付成果应满足双方确认的需求说明、功能清单、数据格式和质量要求。甲方应在收到交付成果后七个工作日内完成验收。",
        "七、保密条款：任何一方对在履行本合同过程中获知的商业秘密、技术资料、客户信息和业务数据负有保密义务，未经对方书面同意不得向第三方披露。",
        "八、违约责任：任何一方违反本合同约定造成对方损失的，应承担相应赔偿责任。因不可抗力导致无法履约的，双方可协商顺延或解除合同。",
        f"九、联系人：甲方指定联系人为{contractData['contactPerson']}，联系电话：{contractData['contactPhone']}。联系人信息变更时，应至少提前三个工作日书面通知对方。",
        f"十、签署日期：{contractData['signDate']}。本合同自双方授权代表签字并加盖公章之日起生效。",
        "（以下无正文）",
        "甲方（盖章）：____________________        乙方（盖章）：____________________",
        "授权代表：________________________        授权代表：________________________",
    ]


# 创建 DOCX 测试文件
def createDocxFile(outputDir, contractData, paragraphList):
    # 第1步：创建 Word 文档对象
    document = Document()

    # 第2步：设置页面边距
    section = document.sections[0]
    section.top_margin = Pt(54)
    section.bottom_margin = Pt(54)
    section.left_margin = Pt(60)
    section.right_margin = Pt(60)

    # 第3步：写入居中标题
    titleParagraph = document.add_paragraph()
    titleParagraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titleRun = titleParagraph.add_run(contractData["title"])
    titleRun.bold = True
    titleRun.font.size = Pt(18)

    # 第4步：写入合同正文段落
    for paragraphText in paragraphList:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Pt(24)
        paragraph.paragraph_format.line_spacing = 1.5
        run = paragraph.add_run(paragraphText)
        run.font.size = Pt(11)

    # 第5步：保存 DOCX 文件
    document.save(outputDir / f"{contractData['fileName']}.docx")


# 创建 PDF 测试文件
def createPdfFile(outputDir, contractData, paragraphList, fontName):
    # 第1步：创建 PDF 画布
    pdfPath = outputDir / f"{contractData['fileName']}.pdf"
    pdf = canvas.Canvas(str(pdfPath), pagesize=A4)
    pageWidth, pageHeight = A4
    currentY = pageHeight - 28 * mm

    # 第2步：写入 PDF 标题
    pdf.setFont(fontName, 18)
    pdf.drawCentredString(pageWidth / 2, currentY, contractData["title"])
    currentY -= 18 * mm

    # 第3步：写入 PDF 正文并自动换行
    pdf.setFont(fontName, 11)
    for paragraphText in paragraphList:
        currentY = drawWrappedText(pdf, paragraphText, fontName, currentY, pageWidth, pageHeight)

    # 第4步：保存 PDF 文件
    pdf.save()


# 绘制自动换行文本
def drawWrappedText(pdf, paragraphText, fontName, currentY, pageWidth, pageHeight):
    # 第1步：初始化行文本和最大宽度
    lineText = ""
    maxWidth = pageWidth - 42 * mm

    # 第2步：逐字计算宽度并换行
    for char in paragraphText:
        candidateText = lineText + char
        if pdf.stringWidth(candidateText, fontName, 11) > maxWidth:
            pdf.drawString(22 * mm, currentY, lineText)
            currentY -= 8 * mm
            lineText = char
            currentY = ensurePdfPage(pdf, currentY, fontName, pageHeight)
        else:
            lineText = candidateText

    # 第3步：绘制剩余文本
    if lineText:
        pdf.drawString(22 * mm, currentY, lineText)
        currentY -= 9 * mm

    # 第4步：检查是否需要新页
    return ensurePdfPage(pdf, currentY, fontName, pageHeight)


# 确保 PDF 当前页还有空间
def ensurePdfPage(pdf, currentY, fontName, pageHeight):
    # 第1步：空间不足时换到新页
    if currentY < 25 * mm:
        pdf.showPage()
        pdf.setFont(fontName, 11)
        return pageHeight - 25 * mm

    # 第2步：空间足够时返回原位置
    return currentY


# 第1步：执行测试文档生成入口
if __name__ == "__main__":
    generateTestDocuments()
